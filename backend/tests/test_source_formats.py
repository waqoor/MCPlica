from datetime import UTC, datetime
from io import BytesIO
from uuid import UUID
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from openpyxl import Workbook

from app.core.exceptions import SourceParseError
from app.domain.sources import ProjectSourceRecord, SourceKind, SourceOrigin
from app.services.sources import _detect_format


def _source() -> ProjectSourceRecord:
    return ProjectSourceRecord(
        id=UUID(int=30),
        project_id=UUID(int=31),
        kind=SourceKind.DOCUMENTATION,
        name="Uploaded documentation",
        origin_type=SourceOrigin.UPLOAD,
        source_url=None,
        is_primary=False,
        created_at=datetime.now(UTC),
    )


def _xlsx() -> bytes:
    output = BytesIO()
    workbook = Workbook()
    workbook.active.append(["GET", "/widgets"])
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _docx() -> bytes:
    output = BytesIO()
    document = Document()
    document.add_paragraph("Widget API")
    document.save(output)
    return output.getvalue()


@pytest.mark.parametrize(
    ("filename", "media_type", "value", "expected"),
    [
        ("guide.json", "application/octet-stream", b'{"api": "widgets"}', "json"),
        ("guide.md", "application/octet-stream", b"# Widgets", "markdown"),
        ("guide.txt", "application/octet-stream", b"Widgets", "text"),
        ("guide.csv", "application/octet-stream", b"method,path\nGET,/widgets", "csv"),
        ("guide.pdf", "application/pdf", b"%PDF-1.7\n", "pdf"),
    ],
)
def test_document_detection_uses_uploaded_filename_and_content(
    filename: str,
    media_type: str,
    value: bytes,
    expected: str,
) -> None:
    assert _detect_format(_source(), value, media_type, filename=filename) == expected


def test_office_packages_are_detected_from_content() -> None:
    assert (
        _detect_format(
            _source(),
            _xlsx(),
            "application/octet-stream",
            filename="guide.xlsx",
        )
        == "xlsx"
    )
    assert (
        _detect_format(
            _source(),
            _docx(),
            "application/octet-stream",
            filename="guide.docx",
        )
        == "docx"
    )


def test_generic_and_macro_enabled_zip_packages_are_rejected() -> None:
    generic = BytesIO()
    with ZipFile(generic, "w", ZIP_DEFLATED) as archive:
        archive.writestr("notes.txt", "not an Office package")
    with pytest.raises(SourceParseError, match="valid XLSX or DOCX"):
        _detect_format(
            _source(),
            generic.getvalue(),
            "application/octet-stream",
            filename="guide.docx",
        )

    macro = BytesIO(_docx())
    with ZipFile(macro, "a", ZIP_DEFLATED) as archive:
        archive.writestr("word/vbaProject.bin", b"macro")
    with pytest.raises(SourceParseError, match="Macro-enabled"):
        _detect_format(
            _source(),
            macro.getvalue(),
            "application/octet-stream",
            filename="guide.docx",
        )
